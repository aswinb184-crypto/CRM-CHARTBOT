import logging
import os
import pickle
import re
import threading
import uuid
from io import BytesIO

import faiss
from docx import Document
from PIL import Image
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

try:
    import pytesseract
except ImportError:
    pytesseract = None


logger = logging.getLogger(__name__)
_index_build_lock = threading.RLock()


# ============================================================
# PATHS
# ============================================================

# ============================================================
# EMBEDDING MODEL
# ============================================================

EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"

print("Loading hospital document embedding model...")

embedding_model = SentenceTransformer(
    EMBEDDING_MODEL_NAME
)

print("Hospital document embedding model loaded.")


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(file_path):
    text_parts = []
    try:
        reader = PdfReader(file_path)
    except Exception as error:
        raise ValueError("Could not open the PDF document.") from error
    pages_needing_ocr = []
    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):
        try:
            page_text = page.extract_text()
        except Exception as error:
            raise ValueError(
                "Could not extract text from PDF page {}.".format(page_number)
            ) from error
        if page_text:
            text_parts.append(
                f"[Page {page_number}]\n{page_text}"
            )
        else:
            pages_needing_ocr.append(page_number - 1)

    if pages_needing_ocr:
        if pdfium is None:
            raise RuntimeError(
                "Scanned PDF OCR is unavailable. Install pypdfium2 and configure "
                "the Tesseract OCR engine."
            )
        _require_ocr_engine()
        pdf_document = pdfium.PdfDocument(file_path)
        try:
            for page_index in pages_needing_ocr:
                page = pdf_document[page_index]
                try:
                    bitmap = page.render(scale=2)
                    try:
                        image = bitmap.to_pil()
                        try:
                            page_text = _ocr_image(image)
                        finally:
                            image.close()
                    finally:
                        bitmap.close()
                finally:
                    page.close()
                if page_text:
                    text_parts.append(
                        "[Page {} - OCR]\n{}".format(page_index + 1, page_text)
                    )
        finally:
            pdf_document.close()

    return "\n\n".join(text_parts)


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(file_path):

    document = Document(file_path)

    paragraphs = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                paragraphs.append(
                    " | ".join(row_text)
                )

    return "\n".join(paragraphs)


# ============================================================
# TXT TEXT EXTRACTION
# ============================================================

def extract_text_from_txt(file_path):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252"
    ]

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding
            ) as file:

                return file.read()

        except UnicodeDecodeError:

            continue

    raise ValueError(
        "Could not read the TXT file encoding."
    )


# ============================================================
# GENERAL TEXT EXTRACTION
# ============================================================

def extract_text(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    elif extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    elif extension == ".txt":

        return extract_text_from_txt(
            file_path
        )

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


def _require_ocr_engine():
    if pytesseract is None:
        raise RuntimeError(
            "OCR is unavailable. Install pytesseract and the Tesseract OCR engine."
        )
    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError as error:
        raise RuntimeError(
            "OCR is unavailable because the Tesseract OCR engine is not installed "
            "or is not on PATH."
        ) from error


def _ocr_image(image):
    _require_ocr_engine()
    try:
        return pytesseract.image_to_string(image).strip()
    except pytesseract.TesseractNotFoundError as error:
        raise RuntimeError(
            "OCR is unavailable because the Tesseract OCR engine is not installed "
            "or is not on PATH."
        ) from error


def extract_text_from_image(image_bytes):
    if pytesseract is None:
        raise RuntimeError(
            "OCR is unavailable. Install pytesseract and the Tesseract OCR engine."
        )

    with Image.open(BytesIO(image_bytes)) as image:
        image.verify()

    with Image.open(BytesIO(image_bytes)) as image:
        image = image.convert("RGB")
        return _ocr_image(image)


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"[ \t]*\n[ \t]*",
        "\n",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# CREATE CHUNKS
# ============================================================

def create_chunks(
    text,
    source_name,
    chunk_size=800,
    overlap=150
):

    text = clean_text(text)

    if not text:
        return []

    chunks = []

    start = 0

    text_length = len(text)

    while start < text_length:

        end = min(
            start + chunk_size,
            text_length
        )

        chunk_text = text[
            start:end
        ].strip()

        if chunk_text:

            chunks.append(
                {
                    "text": chunk_text,
                    "source": source_name
                }
            )

        if end >= text_length:
            break

        start = end - overlap

    return chunks


# ============================================================
# PROCESS ONE DOCUMENT
# ============================================================

def process_document(file_path, source_name=None):

    if not os.path.exists(file_path):
        raise FileNotFoundError("Hospital document is missing.")

    source_name = source_name or os.path.basename(file_path)

    # Extract
    try:
        text = extract_text(file_path)
    except (ModuleNotFoundError, ImportError, RuntimeError, ValueError):
        raise
    except Exception as error:
        raise ValueError("The document is invalid or could not be read.") from error

    # Clean
    text = clean_text(
        text
    )

    if not text:

        raise ValueError(
            "No readable text was found in the document."
        )

    # Create chunks
    chunks = create_chunks(
        text,
        source_name
    )

    if not chunks:

        raise ValueError(
            "Could not create document chunks."
        )

    logger.info(
        "Hospital document extracted; characters=%d chunks=%d",
        len(text),
        len(chunks),
    )
    return chunks


# ============================================================
# BUILD HOSPITAL KNOWLEDGE BASE
# ============================================================

def build_patient_index(file_paths, store_dir, source_names=None):
    with _index_build_lock:
        return _build_patient_index(file_paths, store_dir, source_names)


def _build_patient_index(file_paths, store_dir, source_names=None):

    if not file_paths:

        raise ValueError(
            "No hospital documents were provided."
        )

    all_chunks = []

    # Process all documents
    source_names = source_names or {}
    for file_path in file_paths:

        try:
            document_chunks = process_document(
                file_path,
                source_names.get(os.path.abspath(file_path)),
            )
        except Exception as error:
            logger.exception(
                "Hospital document processing failed; exception_type=%s",
                type(error).__name__,
            )
            raise

        all_chunks.extend(
            document_chunks
        )

    if not all_chunks:

        raise ValueError(
            "No document content was available."
        )

    texts = [
        chunk["text"]
        for chunk in all_chunks
    ]

    embeddings = embedding_model.encode(
        texts,
        convert_to_numpy=True,
        show_progress_bar=True
    )

    embeddings = embeddings.astype(
        "float32"
    )

    dimension = embeddings.shape[1]

    # Create FAISS index
    index = faiss.IndexFlatL2(
        dimension
    )

    index.add(
        embeddings
    )
    logger.info(
        "Hospital document index built; documents=%d chunks=%d",
        len(file_paths),
        len(all_chunks),
    )

    store_dir = os.path.abspath(store_dir)
    os.makedirs(store_dir, exist_ok=True)
    index_file = os.path.join(store_dir, "patient_faiss.index")
    chunks_file = os.path.join(store_dir, "patient_chunks.pkl")
    transaction_id = uuid.uuid4().hex
    temporary_index = index_file + "." + transaction_id + ".tmp"
    temporary_chunks = chunks_file + "." + transaction_id + ".tmp"
    backup_index = index_file + "." + transaction_id + ".bak"
    backup_chunks = chunks_file + "." + transaction_id + ".bak"
    committed = False
    had_index = os.path.exists(index_file)
    had_chunks = os.path.exists(chunks_file)

    try:
        faiss.write_index(index, temporary_index)
        with open(temporary_chunks, "wb") as file:
            pickle.dump(all_chunks, file)
        try:
            if had_index:
                os.replace(index_file, backup_index)
            if had_chunks:
                os.replace(chunks_file, backup_chunks)
            os.replace(temporary_index, index_file)
            os.replace(temporary_chunks, chunks_file)
        except Exception:
            for target, backup, had_original in (
                (index_file, backup_index, had_index),
                (chunks_file, backup_chunks, had_chunks),
            ):
                if os.path.exists(backup):
                    if os.path.exists(target):
                        os.remove(target)
                    os.replace(backup, target)
                elif not had_original and os.path.exists(target):
                    os.remove(target)
            raise
        committed = True
    finally:
        for temporary_path in (
            temporary_index,
            temporary_chunks,
            backup_index if committed else None,
            backup_chunks if committed else None,
        ):
            if temporary_path is None:
                continue
            if os.path.exists(temporary_path):
                try:
                    os.remove(temporary_path)
                except OSError:
                    logger.exception(
                        "Hospital index transaction file cleanup failed."
                    )

    return {
        "documents": len(file_paths),
        "chunks": len(all_chunks),
        "index_file": index_file,
        "chunks_file": chunks_file
    }


# ============================================================
# SEARCH HOSPITAL DOCUMENTS
# ============================================================

def search_patient_documents(
    question,
    top_k=5,
    *,
    store_dir=None,
    hospital_id=None,
):
    """Search the configured hospital index or a trusted explicit store directory."""

    if not isinstance(question, str) or not question.strip():
        raise ValueError("A non-empty search question is required.")
    if isinstance(top_k, bool) or not isinstance(top_k, int) or top_k < 1:
        raise ValueError("top_k must be a positive integer.")
    if store_dir is not None and hospital_id is not None:
        raise ValueError("Specify either store_dir or hospital_id, not both.")

    if store_dir is None:
        from security import hospital_index_dir

        configured_hospital_id = hospital_id
        if configured_hospital_id is None:
            configured_hospital_id = os.environ.get(
                "APP_HOSPITAL_ID",
                "local_hospital",
            )
        store_dir = hospital_index_dir(configured_hospital_id)
    else:
        store_dir = os.fspath(store_dir)
        if not isinstance(store_dir, str):
            raise TypeError("store_dir must be a string path or PathLike object.")

    index_file = os.path.join(store_dir, "patient_faiss.index")
    chunks_file = os.path.join(store_dir, "patient_chunks.pkl")

    if not os.path.exists(index_file):

        raise FileNotFoundError(
            "Hospital document index does not exist yet."
        )

    if not os.path.exists(chunks_file):

        raise FileNotFoundError(
            "Hospital document chunks do not exist yet."
        )

    # Load FAISS
    index = faiss.read_index(
        index_file
    )

    # Load chunks
    with open(
        chunks_file,
        "rb"
    ) as file:

        chunks = pickle.load(
            file
        )

    if not chunks:

        return []

    if index.ntotal != len(chunks):
        raise ValueError("Hospital document index and chunks are inconsistent.")

    # Question embedding
    question_embedding = (
        embedding_model.encode(
            [question],
            convert_to_numpy=True
        )
        .astype("float32")
    )

    actual_top_k = min(
        top_k,
        index.ntotal
    )

    # Search
    distances, indices = index.search(
        question_embedding,
        actual_top_k
    )

    results = []

    for distance, index_number in zip(
        distances[0],
        indices[0]
    ):

        if index_number < 0:
            continue

        chunk = chunks[
            index_number
        ].copy()

        chunk["distance"] = float(
            distance
        )

        results.append(
            chunk
        )

    return results


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    print("=" * 60)

    print(
        "       HOSPITAL DOCUMENT PROCESSOR"
    )

    print("=" * 60)

    print(
        "\nThis module is ready."
    )

    print(
        "\nSupported files:"
    )

    print("PDF")
    print("DOCX")
    print("TXT")

    print(
        "\nThe Streamlit Hospital Knowledge Base "
        "will use this module."
    )