import io
import os
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import httpx
import numpy
from groq import (
    APIConnectionError,
    APITimeoutError,
    AuthenticationError,
    NotFoundError,
    RateLimitError,
)
from PIL import Image, ImageDraw, ImageFont
from docx import Document

import app as app_module
import document.rag as rag
import patient_rag.document_processor as processor
from app import app
from security import hospital_index_dir


class PatientDocumentProcessorTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.previous_environment = {
            name: os.environ.get(name)
            for name in ("HOSPITAL_DATA_DIR", "APP_HOSPITAL_ID")
        }
        os.environ["HOSPITAL_DATA_DIR"] = self.temporary_directory.name
        os.environ["APP_HOSPITAL_ID"] = "processor_test"
        app.config.update(TESTING=True, RATE_LIMIT_REQUESTS=120)
        app_module._request_times.clear()
        self.client = app.test_client()
        self.encoder_patch = patch.object(
            processor.embedding_model,
            "encode",
            side_effect=self._fake_embeddings,
        )
        self.encoder_patch.start()
        self.addCleanup(self.encoder_patch.stop)

    def tearDown(self):
        for name, previous_value in self.previous_environment.items():
            if previous_value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = previous_value
        self.temporary_directory.cleanup()

    @staticmethod
    def _fake_embeddings(texts, **_kwargs):
        return numpy.array(
            [[float(len(text)), 1.0, 0.5] for text in texts],
            dtype="float32",
        )

    @staticmethod
    def _minimal_text_pdf(text):
        escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = (
            "BT /F1 12 Tf 72 720 Td ({}) Tj ET\n".format(escaped_text)
        ).encode("ascii")
        objects = (
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
            ),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            (
                b"<< /Length " + str(len(stream)).encode("ascii")
                + b" >>\nstream\n" + stream + b"endstream"
            ),
        )

        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for object_number, content in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(
                "{} 0 obj\n".format(object_number).encode("ascii")
                + content
                + b"\nendobj\n"
            )

        xref_offset = len(pdf)
        pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend("{:010d} 00000 n \n".format(offset).encode("ascii"))
        pdf.extend(
            b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
            + str(xref_offset).encode("ascii")
            + b"\n%%EOF\n"
        )
        return bytes(pdf)

    def _upload_text(self, filename, content):
        return self.client.post(
            "/documents/upload",
            data={"files": (io.BytesIO(content), filename)},
            content_type="multipart/form-data",
        )

    def test_search_accepts_top_k_positionally_and_defaults_to_configured_index(self):
        source_path = os.path.join(self.temporary_directory.name, "timings.txt")
        with open(source_path, "w", encoding="utf-8") as source_file:
            source_file.write(
                "Hospital outpatient clinic timings are 8 AM to 4 PM."
            )

        store_dir = hospital_index_dir("processor_test")
        processor.build_patient_index([source_path], store_dir=store_dir)

        results = processor.search_patient_documents(
            "hospital outpatient timings",
            5,
        )

        self.assertEqual(len(results), 1)
        self.assertIn("outpatient clinic timings", results[0]["text"])

    def test_search_supports_explicit_keyword_store_dir_and_rejects_bad_top_k(self):
        source_path = os.path.join(self.temporary_directory.name, "policy.txt")
        store_dir = os.path.join(self.temporary_directory.name, "separate-index")
        with open(source_path, "w", encoding="utf-8") as source_file:
            source_file.write("Visitor policy information.")
        processor.build_patient_index([source_path], store_dir=store_dir)

        results = processor.search_patient_documents(
            "visitor policy",
            top_k=2,
            store_dir=store_dir,
        )
        self.assertEqual(len(results), 1)
        with self.assertRaisesRegex(ValueError, "top_k"):
            processor.search_patient_documents("visitor policy", top_k="5")
        with self.assertRaisesRegex(ValueError, "top_k"):
            processor.search_patient_documents("visitor policy", top_k=True)

    def test_text_cleaning_chunk_creation_and_docx_pdf_extraction(self):
        self.assertEqual(
            processor.clean_text("  alpha   beta \n\n\n gamma delta  "),
            "alpha beta\n\ngamma delta",
        )
        chunks = processor.create_chunks(
            "alpha beta\n\ngamma delta",
            "sample.txt",
            chunk_size=12,
            overlap=3,
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk["source"] == "sample.txt" for chunk in chunks))

        docx_path = os.path.join(self.temporary_directory.name, "sample.docx")
        document = Document()
        document.add_paragraph("Outpatient clinic timings are 8 AM to 4 PM.")
        document.save(docx_path)
        self.assertIn(
            "Outpatient clinic timings",
            processor.extract_text_from_docx(docx_path),
        )

        pdf_path = os.path.join(self.temporary_directory.name, "sample.pdf")
        with open(pdf_path, "wb") as pdf_file:
            pdf_file.write(
                self._minimal_text_pdf(
                    "Outpatient clinic timings are 8 AM to 4 PM."
                )
            )
        self.assertIn(
            "Outpatient clinic timings",
            processor.extract_text_from_pdf(pdf_path),
        )

    def test_uploads_pdf_docx_and_txt_and_retrieves_their_contents(self):
        pdf_upload = self._upload_text(
            "hours.pdf",
            self._minimal_text_pdf("Outpatient reception opens at 8 AM."),
        )
        docx_buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("The pharmacy closes at 6 PM.")
        document.save(docx_buffer)
        docx_upload = self._upload_text("pharmacy.docx", docx_buffer.getvalue())
        txt_upload = self._upload_text(
            "visiting.txt",
            b"Visiting hours end at 7 PM.",
        )

        self.assertEqual(pdf_upload.status_code, 201)
        self.assertEqual(docx_upload.status_code, 201)
        self.assertEqual(txt_upload.status_code, 201)
        results = processor.search_patient_documents(
            "pharmacy closing time",
            top_k=10,
        )
        indexed_content = "\n".join(result["text"] for result in results)
        self.assertIn("Outpatient reception opens at 8 AM", indexed_content)
        self.assertIn("The pharmacy closes at 6 PM", indexed_content)
        self.assertIn("Visiting hours end at 7 PM", indexed_content)

    def test_scanned_pdf_pages_use_the_existing_image_ocr_path(self):
        pdf_page = MagicMock()
        pdf_page.extract_text.return_value = ""
        pdf_reader = MagicMock(pages=[pdf_page])
        rendered_page = MagicMock()
        bitmap = MagicMock()
        bitmap.to_pil.return_value = Image.new("RGB", (2, 2), color="white")
        rendered_page.render.return_value = bitmap
        pdf_document = MagicMock()
        pdf_document.__getitem__.return_value = rendered_page

        with tempfile.NamedTemporaryFile(suffix=".pdf") as source_file, patch.object(
            processor,
            "PdfReader",
            return_value=pdf_reader,
        ), patch.object(
            processor.pdfium,
            "PdfDocument",
            return_value=pdf_document,
        ), patch.object(
            processor,
            "_ocr_image",
            return_value="Scanned outpatient hours are 8 AM to 4 PM.",
        ) as ocr_image:
            extracted = processor.extract_text_from_pdf(source_file.name)

        self.assertIn("Scanned outpatient hours are 8 AM to 4 PM", extracted)
        ocr_image.assert_called_once()
        pdf_document.close.assert_called_once()
        rendered_page.close.assert_called_once()
        bitmap.close.assert_called_once()

    def test_scanned_pdf_without_ocr_dependencies_returns_actionable_error(self):
        pdf_page = MagicMock()
        pdf_page.extract_text.return_value = ""
        with tempfile.NamedTemporaryFile(suffix=".pdf") as source_file, patch.object(
            processor,
            "PdfReader",
            return_value=MagicMock(pages=[pdf_page]),
        ), patch.object(processor, "pdfium", None):
            with self.assertRaisesRegex(RuntimeError, "Scanned PDF OCR is unavailable"):
                processor.extract_text_from_pdf(source_file.name)

    def test_image_and_scanned_pdf_are_ocr_indexed_with_local_tesseract(self):
        try:
            processor._require_ocr_engine()
        except RuntimeError as error:
            self.skipTest(str(error))
        if processor.pdfium is None:
            self.skipTest("pypdfium2 is not installed.")

        image = Image.new("RGB", (900, 180), color="white")
        ImageDraw.Draw(image).text(
            (20, 40),
            "Outpatient clinic hours are 8 AM to 4 PM",
            fill="black",
            font=ImageFont.load_default(size=42),
        )
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")

        pdf_buffer = io.BytesIO()
        image.save(pdf_buffer, format="PDF", resolution=150)
        image_upload = self.client.post(
            "/documents/upload",
            data={"files": (io.BytesIO(image_buffer.getvalue()), "hours.png")},
            content_type="multipart/form-data",
        )
        pdf_upload = self._upload_text("scanned-hours.pdf", pdf_buffer.getvalue())

        self.assertEqual(image_upload.status_code, 201)
        self.assertEqual(pdf_upload.status_code, 201)
        results = processor.search_patient_documents(
            "outpatient clinic hours",
            top_k=10,
        )
        extracted_content = "\n".join(result["text"] for result in results)
        self.assertIn("Outpatient clinic hours are 8 AM to 4 PM", extracted_content)
        self.assertGreaterEqual(len(results), 2)

    def test_failed_index_publication_restores_the_previous_index_and_metadata(self):
        index_dir = os.path.join(self.temporary_directory.name, "atomic-index")
        original_path = os.path.join(self.temporary_directory.name, "original.txt")
        replacement_path = os.path.join(self.temporary_directory.name, "replacement.txt")
        with open(original_path, "w", encoding="utf-8") as source_file:
            source_file.write("Original outpatient hours are 8 AM to 4 PM.")
        with open(replacement_path, "w", encoding="utf-8") as source_file:
            source_file.write("Replacement emergency hours are 24 hours.")

        processor.build_patient_index([original_path], index_dir)
        index_path = os.path.join(index_dir, "patient_faiss.index")
        chunks_path = os.path.join(index_dir, "patient_chunks.pkl")
        with open(index_path, "rb") as file:
            original_index = file.read()
        with open(chunks_path, "rb") as file:
            original_chunks = file.read()
        real_replace = os.replace

        def fail_chunks_publish(source, destination):
            if destination == chunks_path and source.endswith(".tmp"):
                raise OSError("simulated metadata publication failure")
            return real_replace(source, destination)

        with patch.object(processor.os, "replace", side_effect=fail_chunks_publish):
            with self.assertRaisesRegex(OSError, "metadata publication failure"):
                processor.build_patient_index([replacement_path], index_dir)

        with open(index_path, "rb") as file:
            self.assertEqual(file.read(), original_index)
        with open(chunks_path, "rb") as file:
            self.assertEqual(file.read(), original_chunks)
        results = processor.search_patient_documents(
            "outpatient hours",
            top_k=5,
            store_dir=index_dir,
        )
        self.assertIn("Original outpatient hours", results[0]["text"])

    def test_upload_index_individual_removal_and_remove_all_use_current_documents(self):
        first_upload = self._upload_text(
            "outpatient.txt",
            b"Outpatient clinic timings are 8 AM to 4 PM.",
        )
        second_upload = self._upload_text(
            "emergency.txt",
            b"Emergency department entrance is on the east side.",
        )
        self.assertEqual(first_upload.status_code, 201)
        self.assertEqual(second_upload.status_code, 201)

        indexed_results = processor.search_patient_documents(
            "hospital outpatient timings",
            top_k=5,
        )
        indexed_text = "\n".join(result["text"] for result in indexed_results)
        self.assertIn("Outpatient clinic timings", indexed_text)
        self.assertIn("Emergency department entrance", indexed_text)

        first_document_id = first_upload.get_json()["documents"][0]["id"]
        removed = self.client.delete("/documents/" + first_document_id)
        self.assertEqual(removed.status_code, 200)
        after_individual_removal = processor.search_patient_documents(
            "hospital outpatient timings",
            top_k=5,
        )
        remaining_text = "\n".join(
            result["text"] for result in after_individual_removal
        )
        self.assertNotIn("Outpatient clinic timings", remaining_text)
        self.assertIn("Emergency department entrance", remaining_text)

        removed_all = self.client.delete("/documents")
        self.assertEqual(removed_all.status_code, 200)
        self.assertEqual(self.client.get("/documents").get_json()["documents"], [])
        index_dir = hospital_index_dir("processor_test")
        self.assertFalse(os.path.exists(os.path.join(index_dir, "patient_faiss.index")))
        self.assertFalse(os.path.exists(os.path.join(index_dir, "patient_chunks.pkl")))

    def test_image_ocr_text_is_indexed_without_sending_binary_image_to_text_extractor(self):
        image_buffer = io.BytesIO()
        Image.new("RGB", (2, 2), color="white").save(image_buffer, format="PNG")

        with patch.object(
            processor,
            "extract_text_from_image",
            return_value="Outpatient check-in begins at 7:30 AM.",
        ):
            uploaded = self.client.post(
                "/documents/upload",
                data={
                    "files": (
                        io.BytesIO(image_buffer.getvalue()),
                        "clinic.png",
                    )
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(uploaded.status_code, 201)
        results = processor.search_patient_documents(
            "clinic check-in",
            top_k=5,
        )
        self.assertEqual(len(results), 1)
        self.assertIn("check-in begins at 7:30 AM", results[0]["text"])

    def test_image_ocr_unavailable_returns_safe_error_and_keeps_existing_index(self):
        existing = self._upload_text(
            "policy.txt",
            b"Existing outpatient clinic hours are 8 AM to 4 PM.",
        )
        self.assertEqual(existing.status_code, 201)
        image_buffer = io.BytesIO()
        Image.new("RGB", (2, 2), color="white").save(image_buffer, format="PNG")

        with patch.object(
            processor,
            "extract_text_from_image",
            side_effect=RuntimeError("OCR is unavailable on this server."),
        ):
            response = self.client.post(
                "/documents/upload",
                data={"files": (io.BytesIO(image_buffer.getvalue()), "scan.png")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 422)
        self.assertIn("OCR is unavailable", response.get_json()["message"])
        self.assertNotIn("server.", response.get_data(as_text=True))
        self.assertEqual(len(self.client.get("/documents").get_json()["documents"]), 1)
        self.assertIn(
            "Existing outpatient clinic hours",
            processor.search_patient_documents("outpatient clinic hours")[0]["text"],
        )

    def test_upload_index_failure_returns_safe_error_and_removes_partial_upload(self):
        with patch(
            "app._index_hospital_documents",
            side_effect=RuntimeError("private internal indexing failure"),
        ):
            response = self._upload_text(
                "failed.txt",
                b"Temporary document content.",
            )

        self.assertEqual(response.status_code, 500)
        self.assertEqual(
            response.get_json(),
            {
                "status": "error",
                "message": (
                    "The uploaded documents could not be processed. Review the local "
                    "server log and try again."
                ),
            },
        )
        self.assertEqual(self.client.get("/documents").get_json()["documents"], [])
        documents_dir = os.path.join(
            self.temporary_directory.name,
            "processor_test",
            "documents",
        )
        self.assertEqual(os.listdir(documents_dir), ["manifest.json"])

    def test_failed_upload_preserves_previously_indexed_documents(self):
        previous = self._upload_text(
            "previous.txt",
            b"Outpatient services operate from 8 AM to 4 PM.",
        )
        self.assertEqual(previous.status_code, 201)

        with patch(
            "app._index_hospital_documents",
            side_effect=RuntimeError("simulated index failure"),
        ):
            failed = self._upload_text(
                "new.txt",
                b"New emergency department information.",
            )

        self.assertEqual(failed.status_code, 500)
        self.assertEqual(len(self.client.get("/documents").get_json()["documents"]), 1)
        search_results = processor.search_patient_documents(
            "outpatient services",
            top_k=5,
        )
        self.assertTrue(search_results)
        self.assertIn("Outpatient services operate", search_results[0]["text"])

    def test_failed_reindex_preserves_the_previous_working_index(self):
        upload = self._upload_text(
            "policy.txt",
            b"Outpatient clinic timings are 8 AM to 4 PM.",
        )
        self.assertEqual(upload.status_code, 201)

        with patch.object(
            processor,
            "build_patient_index",
            side_effect=RuntimeError("internal indexing failure"),
        ):
            response = self.client.post("/documents/reindex")

        self.assertEqual(response.status_code, 500)
        results = processor.search_patient_documents(
            "outpatient clinic timings",
            top_k=5,
        )
        self.assertTrue(results)
        self.assertIn("Outpatient clinic timings", results[0]["text"])

    def test_uploaded_document_summary_phrases_route_actual_chunks_to_summary(self):
        uploaded_chunks = [
            {
                "text": "Outpatient services operate from 8 AM to 4 PM.",
                "source": "outpatient-policy.pdf",
                "index": 0,
            }
        ]
        questions = (
            "summarize this file",
            "tell me about the PDF I uploaded",
            "give me about it",
            "give me summary about it",
            "give me a summary",
        )

        for question in questions:
            with self.subTest(question=question):
                self.assertTrue(rag.is_document_summary_question(question))
                with patch.object(
                    rag,
                    "load_all_hospital_chunks",
                    return_value=uploaded_chunks,
                ) as load_chunks, patch.object(
                    rag,
                    "summarize_hospital_document",
                    return_value="The uploaded policy describes outpatient hours.",
                ) as summarize:
                    answer = rag.answer_question(
                        question,
                        hospital_id="processor_test",
                    )

                self.assertEqual(
                    answer,
                    "The uploaded policy describes outpatient hours.",
                )
                load_chunks.assert_called_once_with("processor_test")
                summarize.assert_called_once_with(question, uploaded_chunks)

    def test_duplicate_uploads_are_indexed_once_and_removal_rebuilds_index(self):
        content = b"Outpatient timings are Monday through Saturday, 9 AM to 5 PM."
        first_upload = self._upload_text("policy.txt", content)
        second_upload = self._upload_text("policy-copy.txt", content)
        self.assertEqual(first_upload.status_code, 201)
        self.assertEqual(second_upload.status_code, 200)
        self.assertEqual(second_upload.get_json()["documents"], [])
        self.assertEqual(len(second_upload.get_json()["duplicates"]), 1)
        self.assertEqual(len(self.client.get("/documents").get_json()["documents"]), 1)

        results = processor.search_patient_documents(
            "outpatient timings",
            top_k=10,
        )
        self.assertEqual(len(results), 1)
        self.assertIn("9 AM to 5 PM", results[0]["text"])
        with patch.object(
            rag.embedding_model,
            "encode",
            side_effect=self._fake_embeddings,
        ):
            self.assertEqual(
                len(
                    rag.search_hospital_documents(
                        "outpatient timings",
                        "processor_test",
                        top_k=10,
                        threshold=1_000_000,
                    )
                ),
                1,
            )

        unique_upload = self._upload_text(
            "emergency.txt",
            b"Emergency entrance is on the east side.",
        )
        self.assertEqual(unique_upload.status_code, 201)

        first_id = first_upload.get_json()["documents"][0]["id"]
        self.assertEqual(self.client.delete("/documents/" + first_id).status_code, 200)
        remaining_results = processor.search_patient_documents(
            "emergency entrance",
            top_k=10,
        )
        self.assertEqual(len(remaining_results), 1)
        self.assertIn("Emergency entrance", remaining_results[0]["text"])
        with patch.object(
            rag.embedding_model,
            "encode",
            side_effect=self._fake_embeddings,
        ):
            self.assertEqual(
                len(
                    rag.search_hospital_documents(
                        "outpatient timings",
                        "processor_test",
                        top_k=10,
                        threshold=1_000_000,
                    )
                ),
                1,
            )

        unique_id = unique_upload.get_json()["documents"][0]["id"]
        self.assertEqual(self.client.delete("/documents/" + unique_id).status_code, 200)
        with self.assertRaises(FileNotFoundError):
            processor.search_patient_documents(
                "outpatient timings",
                top_k=10,
            )
        self.assertEqual(
            rag.search_hospital_documents(
                "outpatient timings",
                "processor_test",
                top_k=10,
                threshold=1_000_000,
            ),
            [],
        )

    def test_hospital_answers_and_summary_receive_indexed_pdf_text(self):
        document_text = (
            "Outpatient services run Monday to Saturday from 9:00 AM to 5:00 PM. "
            "The emergency department operates 24 hours a day, 7 days a week. "
            "The reception phone number is 0422-555-0100."
        )
        zero_embeddings = lambda texts, **_kwargs: numpy.zeros(
            (len(texts), 3),
            dtype="float32",
        )
        with patch.object(
            processor.embedding_model,
            "encode",
            side_effect=zero_embeddings,
        ), patch.object(
            rag.embedding_model,
            "encode",
            side_effect=zero_embeddings,
        ):
            upload = self._upload_text(
                "Hospital_Test_Knowledge_Base.pdf",
                self._minimal_text_pdf(document_text),
            )
            self.assertEqual(upload.status_code, 201)

            cases = (
                ("What are the outpatient timings?", "9:00 AM to 5:00 PM"),
                (
                    "What are the emergency department timings?",
                    "24 hours a day, 7 days a week",
                ),
                ("What is the reception phone number?", "0422-555-0100"),
            )
            for question, expected_fact in cases:
                with self.subTest(question=question), patch.object(
                    rag,
                    "generate_groq_answer",
                    return_value="The indexed PDF answers this question.",
                ) as generate_answer:
                    answer = rag.answer_question(
                        question,
                        hospital_id="processor_test",
                    )
                    self.assertEqual(
                        answer,
                        "The indexed PDF answers this question.",
                    )
                    prompt = generate_answer.call_args.args[0]
                    self.assertIn(expected_fact, prompt)

            with patch.object(
                rag,
                "generate_groq_answer",
                return_value="The document summary.",
            ) as generate_answer:
                summary = rag.answer_question(
                    "Give me a summary of the file I uploaded.",
                    hospital_id="processor_test",
                )
            self.assertEqual(summary, "The document summary.")
            self.assertIn(
                "9:00 AM to 5:00 PM",
                generate_answer.call_args.args[0],
            )

    def test_uploaded_pdf_precedes_crm_route_and_supplies_groq_context(self):
        self.encoder_patch.stop()
        synthetic_fact = (
            "The fictional Physiotherapy department opens at 8:40 AM. "
            "Its unique verification phrase is NEBULA-731."
        )
        upload = self._upload_text(
            "Hospital_Test_Knowledge_Base (1).pdf",
            self._minimal_text_pdf(synthetic_fact),
        )
        self.assertEqual(upload.status_code, 201)

        hospital_index, hospital_chunks = rag.load_hospital_knowledge_base(
            "processor_test"
        )
        self.assertIsNotNone(hospital_index)
        self.assertEqual(hospital_index.ntotal, len(hospital_chunks))
        self.assertEqual(
            rag.get_chunk_source(hospital_chunks[0]),
            "Hospital_Test_Knowledge_Base (1).pdf",
        )
        retrieved = rag.search_hospital_documents(
            "According to the uploaded PDF, what time does physiotherapy open?",
            "processor_test",
            top_k=5,
            threshold=1.3,
        )
        self.assertTrue(retrieved)
        self.assertEqual(
            retrieved[0]["source"],
            "Hospital_Test_Knowledge_Base (1).pdf",
        )
        self.assertIn("8:40 AM", retrieved[0]["text"])

        with patch.object(
            rag,
            "generate_groq_answer",
            return_value="The Physiotherapy department opens at 8:40 AM.",
        ) as generate_answer, patch.object(
            rag,
            "search_crm_documents",
        ) as search_crm:
            response = self.client.post(
                "/chat",
                json={
                    "question": (
                        "According to the uploaded PDF, what time does "
                        "physiotherapy open?"
                    )
                },
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["answer"],
            "The Physiotherapy department opens at 8:40 AM.",
        )
        prompt = generate_answer.call_args.args[0]
        self.assertIn("Hospital_Test_Knowledge_Base (1).pdf", prompt)
        self.assertIn("8:40 AM", prompt)
        search_crm.assert_not_called()

    def test_unsupported_uploaded_pdf_question_does_not_get_fabricated_answer(self):
        self.encoder_patch.stop()
        upload = self._upload_text(
            "Hospital_Test_Knowledge_Base.pdf",
            self._minimal_text_pdf(
                "The fictional clinic opens at 8:40 AM. Its phone extension is 4829."
            ),
        )
        self.assertEqual(upload.status_code, 201)

        with patch.object(rag, "generate_groq_answer") as generate_answer:
            response = self.client.post(
                "/chat",
                json={
                    "question": (
                        "According to the uploaded PDF, when is the lunar eclipse "
                        "at the observatory?"
                    )
                },
            )

        self.assertEqual(response.status_code, 200)
        answer = response.get_json()["answer"]
        self.assertIn("couldn't find", answer.lower())
        self.assertNotIn("eclipse", answer.lower())
        generate_answer.assert_not_called()

    def test_explicit_document_search_uses_lexical_evidence_for_borderline_scores(self):
        relevant = {
            "text": "The hospital departments include cardiology and neurology.",
            "source": "hospital.pdf",
            "distance": 1.57,
            "index": 0,
        }
        unrelated = {
            "text": "Outpatient hours are Monday through Saturday.",
            "source": "hospital.pdf",
            "distance": 1.59,
            "index": 1,
        }
        question = "Which departments does the uploaded PDF mention?"

        with patch.object(
            rag,
            "search_hospital_documents",
            return_value=[relevant, unrelated],
        ) as search_hospital, patch.object(
            rag,
            "answer_from_hospital_documents",
            return_value="The document lists cardiology and neurology.",
        ) as answer_from_document:
            answer = rag.answer_question(
                question,
                hospital_id="processor_test",
            )

        self.assertEqual(
            search_hospital.call_args.kwargs["threshold"],
            rag.EXPLICIT_DOCUMENT_DISTANCE,
        )
        answer_from_document.assert_called_once_with(question, [relevant])
        self.assertIn("cardiology", answer)

    def test_groq_authentication_and_model_failures_are_actionable(self):
        fake_response_401 = httpx.Response(
            401,
            request=httpx.Request("POST", "https://api.groq.com/openai/v1/chat/completions"),
        )
        fake_response_404 = httpx.Response(
            404,
            request=httpx.Request("POST", "https://api.groq.com/openai/v1/chat/completions"),
        )
        cases = (
            (
                AuthenticationError(
                    "invalid credentials",
                    response=fake_response_401,
                    body={"error": {"message": "invalid credentials"}},
                ),
                "Groq rejected its API credentials",
            ),
            (
                NotFoundError(
                    "model unavailable",
                    response=fake_response_404,
                    body={"error": {"message": "model not found"}},
                ),
                "could not use the configured model",
            ),
            (
                RateLimitError(
                    "rate limit",
                    response=httpx.Response(
                        429,
                        request=httpx.Request(
                            "POST",
                            "https://api.groq.com/openai/v1/chat/completions",
                        ),
                    ),
                    body={"error": {"message": "rate limit"}},
                ),
                "rate-limiting",
            ),
            (
                APITimeoutError(
                    request=httpx.Request(
                        "POST",
                        "https://api.groq.com/openai/v1/chat/completions",
                    ),
                ),
                "did not respond in time",
            ),
            (
                APIConnectionError(
                    message="connection unavailable",
                    request=httpx.Request(
                        "POST",
                        "https://api.groq.com/openai/v1/chat/completions",
                    ),
                ),
                "could not reach Groq",
            ),
        )

        mocked_client = MagicMock()
        with patch.object(rag, "groq_client", mocked_client), patch.dict(
            os.environ,
            {"GROQ_API_KEY": "test-groq-key-never-used-for-network"},
        ):
            for exception, expected_message in cases:
                with self.subTest(exception=type(exception).__name__):
                    mocked_client.chat.completions.create.side_effect = exception
                    with self.assertRaises(rag.GroqServiceError) as raised:
                        rag.generate_groq_answer("A synthetic test prompt.")
                    self.assertIn(
                        expected_message,
                        raised.exception.user_message,
                    )
                    self.assertNotIn(
                        os.environ.get("GROQ_API_KEY", "no-configured-key"),
                        raised.exception.user_message,
                    )
                    mocked_client.chat.completions.create.side_effect = exception
                    with patch.object(
                        app_module,
                        "answer_question",
                        side_effect=lambda *_args: rag.generate_groq_answer(
                            "A synthetic test prompt."
                        ),
                    ):
                        response = self.client.post(
                            "/chat",
                            json={"question": "What does the test document say?"},
                        )
                    self.assertEqual(
                        response.status_code,
                        raised.exception.http_status,
                    )
                    self.assertIn(
                        expected_message,
                        response.get_json()["message"],
                    )
                    self.assertNotIn(
                        "test-groq-key-never-used-for-network",
                        response.get_data(as_text=True),
                    )

    def test_hospital_question_without_matching_documents_does_not_use_general_answers(self):
        with patch.object(
            rag,
            "search_hospital_documents",
            return_value=[],
        ) as search_hospital, patch.object(
            rag,
            "search_crm_documents",
        ) as search_crm, patch.object(
            rag,
            "generate_groq_answer",
        ) as generate_answer:
            answer = rag.answer_question(
                "What is the hospital's chief surgeon's name?",
                hospital_id="processor_test",
            )

        self.assertIn(
            "couldn't find that information in",
            answer.lower(),
        )
        search_hospital.assert_called_once()
        search_crm.assert_not_called()
        generate_answer.assert_not_called()

    def test_summary_prompt_contains_uploaded_document_content(self):
        chunk = {
            "text": "Outpatient services operate from 8 AM to 4 PM.",
            "source": "outpatient-policy.pdf",
            "index": 0,
        }
        with patch.object(
            rag,
            "generate_groq_answer",
            return_value="The document describes outpatient hours.",
        ) as generate_answer:
            summary = rag.summarize_hospital_document(
                "summarize this file",
                [chunk],
            )

        self.assertEqual(summary, "The document describes outpatient hours.")
        prompt = generate_answer.call_args.args[0]
        self.assertIn(chunk["text"], prompt)
        self.assertIn(chunk["source"], prompt)


if __name__ == "__main__":
    unittest.main()
